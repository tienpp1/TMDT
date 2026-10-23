"""Rebuild deterministic diagrams and coursework DOCX using bundled Python.

Sources in Downloads are read only. Runtime test summary may be replaced in
test-summary.txt only after evidence has been collected.
"""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from hashlib import sha256
import json, math, sys, html
import xml.etree.ElementTree as ET
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')
HERE = Path(__file__).parent
FIG = HERE / 'figures'
FIG.mkdir(exist_ok=True)
SOURCE = Path(r'C:\Users\anhbe\Downloads\TH4 Bổ sung thiết kế cho đồ án cuối kỳ.docx')
FONT = Path(r'C:\Windows\Fonts\arial.ttf')
BOLD = Path(r'C:\Windows\Fonts\arialbd.ttf')
INK, BLUE, PALE, GRAY = '#13283d', '#245a78', '#eaf2f6', '#687987'

class Figure:
    def __init__(self, name, w, h):
        self.name, self.w, self.h = name, w, h
        self.im = Image.new('RGB', (w*2,h*2), 'white')
        self.draw = ImageDraw.Draw(self.im)
        self.svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}">', '<rect width="100%" height="100%" fill="white"/>']
    def line(self, points, color=BLUE, width=3, dash=False):
        pts = [(int(x*2),int(y*2)) for x,y in points]
        if dash:
            for (ax,ay),(bx,by) in zip(pts,pts[1:]):
                length=math.hypot(bx-ax,by-ay)
                for start in range(0,int(length),30):
                    end=min(start+18,length)
                    self.draw.line([(ax+(bx-ax)*start/length,ay+(by-ay)*start/length),(ax+(bx-ax)*end/length,ay+(by-ay)*end/length)],fill=color,width=width*2)
        else:
            self.draw.line(pts,fill=color,width=width*2)
        self.svg.append(f'<polyline points="'+ ' '.join(f'{x},{y}' for x,y in points)+f'" fill="none" stroke="{color}" stroke-width="{width}"'+(' stroke-dasharray="8 6"' if dash else '')+'/>')
    def arrow(self, x1,y1,x2,y2, color=BLUE, dash=False):
        self.line([(x1,y1),(x2,y2)], color, 3, dash)
        a = math.atan2(y2-y1,x2-x1)
        points=[(x2-13*math.cos(a-.45),y2-13*math.sin(a-.45)),(x2,y2),(x2-13*math.cos(a+.45),y2-13*math.sin(a+.45))]
        self.line(points,color,3)
    def rect(self,x,y,w,h,fill=PALE,outline=BLUE):
        self.draw.rectangle((x*2,y*2,(x+w)*2,(y+h)*2),fill=fill,outline=outline,width=4)
        self.svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="3" fill="{fill}" stroke="{outline}" stroke-width="2"/>')
    def text(self,x,y,s,size=25,bold=False,center=False,color=INK):
        font=ImageFont.truetype(str(BOLD if bold else FONT),size*2)
        self.draw.text((x*2,y*2),s,font=font,fill=color,anchor='mt' if center else 'lt')
        self.svg.append(f'<text x="{x}" y="{y+size*.83}" font-family="Arial" font-size="{size}"'+(' font-weight="bold"' if bold else '')+(' text-anchor="middle"' if center else '')+f' fill="{color}">{html.escape(s)}</text>')
    def box(self,x,y,w,h,title,lines,size=24):
        self.rect(x,y,w,h)
        self.text(x+w/2,y+17,title,size+2,True,True)
        for i,s in enumerate(lines): self.text(x+w/2,y+58+i*(size+9),s,size,False,True)
    def save(self):
        self.svg.append('</svg>')
        (FIG/(self.name+'.svg')).write_text('\n'.join(self.svg),encoding='utf-8')
        self.im.save(FIG/(self.name+'.png'),optimize=True)

f=Figure('architecture',1420,540)
f.box(25,35,315,150,'SIMULATOR ĐỘC LẬP',['3 điểm đo giả lập','Chu kỳ gửi 5 giây'])
f.box(520,35,350,150,'EXPRESS API',['Khóa / phiên / quyền','Validate và transaction'])
f.box(1050,35,340,150,'SQLITE',['Readings, rules, alerts','State và audit'])
f.arrow(340,110,520,110); f.text(430,65,'HTTP JSON',21,center=True)
f.arrow(870,110,1050,110); f.text(960,65,'Lưu bền vững',21,center=True)
f.box(520,300,350,150,'LOGIC DÙNG CHUNG',['3 mẫu; 10 / 20 / 7 cm','Freshness, retry, gap'])
f.arrow(695,185,695,300)
f.arrow(870,370,1220,185)
f.box(1050,300,340,150,'WEBSITE',['Manager và staff','Poll 5 giây / SSE'])
f.arrow(1050,320,870,150); f.text(1075,238,'API có quyền',21,center=True)
f.box(25,300,315,150,'PHẦN CỨNG DỰ KIẾN',['Siêu âm + ESP32 Wi-Fi','Chưa chế tạo / đo thật'])
f.arrow(340,360,520,165,dash=True)
f.text(710,498,'Prototype chạy cục bộ; simulator thay thiết bị, không ghi thẳng DB.',24,center=True)
f.save()

f=Figure('erd',1420,650)
for args in [
 (25,25,320,150,'users',['id; email; password_hash','display_name; is_active']),
 (550,25,320,150,'site_users',['site_id + user_id (PK)','role: manager / staff']),
 (1075,25,320,150,'sites',['id; name; location','is_active']),
 (550,250,320,170,'devices',['id; site_id; key_hash','latest_*; *_count','current_alert_level']),
 (25,465,320,150,'readings',['device_id + message_id','time; level; quality','history_only']),
 (550,465,320,150,'alert_rules',['device_id; 10 / 20 / 7','sample_count; version']),
 (1075,465,320,150,'alerts',['device_id; episode_id','ack_by; ack_at','resolved_at; rule_version'])]: f.box(*args,size=22)
f.arrow(345,100,550,100); f.text(450,57,'1 → N',22,center=True)
f.arrow(1075,100,870,100); f.text(972,57,'1 → N',22,center=True)
f.arrow(1235,175,870,300); f.text(1100,228,'1 → N',22,center=True)
f.arrow(550,360,345,465); f.text(420,385,'1 → N',22,center=True)
f.arrow(710,420,710,465)
f.arrow(870,365,1075,465); f.text(1000,385,'1 → N',22,center=True)
f.text(200,290,'Unique reading:',24,bold=True,center=True)
f.text(200,331,'(device_id, message_id)',24,center=True)
f.text(1220,300,'Cùng một đợt:',24,bold=True,center=True)
f.text(1220,342,'ack ≠ resolve',25,center=True)
f.save()

def seq(name,labels,events,height=600):
    f=Figure(name,1420,height)
    xs=[80+i*(1260/(len(labels)-1)) for i in range(len(labels))]
    for x,label in zip(xs,labels):
        f.rect(x-75,20,150,53)
        f.text(x,32,label,22,True,True)
        f.line([(x,74),(x,height-25)],'#a8b3bc',2,True)
    for k,(a,b,label) in enumerate(events):
        y=122+k*65
        if a==b:
            f.line([(xs[a],y),(xs[a]+45,y),(xs[a]+45,y+25),(xs[a],y+25)])
            f.text(xs[a]+53,y-8,label,20)
        else:
            f.arrow(xs[a],y,xs[b],y)
            f.text((xs[a]+xs[b])/2,y-30,label,22,center=True)
    f.save()
seq('sequence-telemetry',['Simulator','API','Logic','SQLite','Website'],[
 (0,1,'POST số đo + khóa thiết bị'),(1,3,'BEGIN, kiểm tra ID và độ mới'),
 (1,2,'Mẫu tươi + trạng thái + luật'),(2,1,'Bộ đếm / chuyển mức'),
 (1,3,'Lưu reading + state + alert; COMMIT'),(1,0,'201 mới / 200 duplicate'),
 (1,4,'SSE báo thay đổi'),(4,1,'GET theo quyền; poll 5 giây')],height=670)
seq('sequence-ack',['Người trực','Website','API','SQLite','Simulator'],[
 (0,1,'Chọn xác nhận tiếp nhận'),(1,2,'POST ack + phiên đăng nhập'),
 (2,3,'Kiểm tra quyền; ghi ack và audit'),(2,1,'Đã tiếp nhận; chưa kết thúc'),
 (4,2,'Ba mẫu tươi liên tiếp <7 cm'),(2,3,'resolved_at; đóng đợt; normal'),
 (2,1,'Thông báo cập nhật')],height=600)

f=Figure('state-alerts',1420,460)
f.box(30,120,330,120,'NORMAL',['Bình thường'],26)
f.box(545,120,330,120,'WARNING',['Cảnh báo'],26)
f.box(1060,120,330,120,'CRITICAL',['Nguy hiểm'],26)
f.arrow(360,165,545,165); f.text(452,95,'3 mẫu ≥10',24,center=True)
f.arrow(875,165,1060,165); f.text(967,95,'3 mẫu ≥20',24,center=True)
f.line([(195,120),(195,45),(1225,45)])
f.arrow(1225,45,1225,120)
f.text(710,5,'NORMAL → CRITICAL trực tiếp khi đủ 3 mẫu ≥20 cm',23,center=True)
f.line([(1225,240),(1225,315),(195,315)])
f.line([(710,240),(710,315)])
f.arrow(195,315,195,240); f.text(710,325,'WARNING hoặc CRITICAL → NORMAL: 3 mẫu <7 cm',24,center=True)
f.text(710,365,'Mẫu lỗi / gap >30 giây: reset đếm. Mẫu cũ / retry: không tăng đếm.',25,center=True)
f.text(710,408,'Ack không đóng cảnh báo. Offline / stale không tự trả normal.',25,center=True)
f.save()

if '--figures-only' in sys.argv:
    print('Created five SVG/PNG figures only; DOCX files unchanged.')
    raise SystemExit(0)

def normal_doc():
    d=Document(); sec=d.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.5); sec.bottom_margin=Cm(1.5)
    sec.left_margin=Cm(1.7); sec.right_margin=Cm(1.7)
    sec.header_distance=Cm(.5); sec.footer_distance=Cm(.5)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3'):
        s=d.styles[name]; s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0)
        s.font.size=Pt(10.5 if name=='Normal' else 13)
        s.paragraph_format.space_after=Pt(5)
        s.paragraph_format.line_spacing=1.05
    for border in d.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    d.styles['Title'].font.size=Pt(17); d.styles['Title'].font.bold=True
    d.styles['Heading 1'].font.size=Pt(13); d.styles['Heading 1'].font.bold=True
    d.styles['Heading 2'].font.size=Pt(11); d.styles['Heading 2'].font.bold=True
    p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=p.add_run('Nhóm 10  •  '); r.font.size=Pt(8)
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); p._p.append(field)
    return d

def para(d,s,bold=False):
    p=d.add_paragraph(); r=p.add_run(s); r.bold=bold; return p

def heading(d,s): d.add_heading(s,level=1)
def image(d,name,width=17.3):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG/(name+'.png')),width=Cm(width))
    p.paragraph_format.space_after=Pt(3)
def caption(d,s):
    p=para(d,s); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: r.font.size=Pt(8)

def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for j,s in enumerate(headers): t.rows[0].cells[j].text=s
    for row in rows:
        cells=t.add_row().cells
        for j,s in enumerate(row): cells[j].text=str(s)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        if i==0:
            repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
        for j,c in enumerate(row.cells):
            c.width=Cm(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'5'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','bottom','left','right'):
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'75'); e.set(qn('w:type'),'dxa'); margins.append(e)
            tcpr.append(margins)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EDF2' if i==0 else ('F5F7F9' if i%2==0 else 'FFFFFF')); tcpr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.0
                for r in p.runs: r.font.name='Arial'; r.font.size=Pt(9); r.bold=i==0
    d.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

d=normal_doc()
d.add_paragraph('Thiết kế và mô phỏng hệ thống IoT giám sát mực nước và cảnh báo ngập cho bãi giữ xe',style='Title')
para(d,'Học phần Thương mại điện tử  |  Hướng 2  |  Nhóm 10  |  23/10/2026',True)
para(d,'Hà Minh Tiến 1250080194 • Trương Quốc Toàn 1250080203 • Huỳnh Nguyễn Duy Tiến 1250080195')
heading(d,'1 Bài toán và phạm vi sản phẩm')
para(d,'Người vận hành bãi xe khó theo dõi liên tục các điểm trũng khi nước dâng. Giải pháp giúp chủ bãi/quản lý và nhân viên trực xem mực nước, cảnh báo vượt ngưỡng, xác nhận tiếp nhận và tra lịch sử để phối hợp ứng phó theo quy trình cơ sở.')
para(d,'Prototype gồm một bãi xe, ba điểm đo giả lập, simulator độc lập gửi HTTP thật, backend Node.js/Express, SQLite và website HTML/CSS/JavaScript. Dữ liệu đo đều mô phỏng; không yêu cầu mua thiết bị, chưa thử thực địa và chưa chứng minh độ chính xác hay hiệu quả giảm thiệt hại.')
heading(d,'Cấu tạo và nguyên lý dự kiến')
para(d,'Thiết bị đề xuất: đầu dò siêu âm, ESP32 Wi-Fi, nguồn DC, hộp bảo vệ và giá đỡ. Đặt cảm biến cao H=100 cm so với mốc đáy; khoảng cách đến mặt nước d=10–100 cm cho h=H−d=0–90 cm. Đo/gửi dự kiến mỗi 5 giây. Lỗi đo dùng quality=sensor_error và level_cm=null. Đây là giả định thiết kế, cần đối chiếu thước đo, vị trí lắp và điều kiện môi trường khi có phần cứng.')
image(d,'architecture')
caption(d,'Hình 1  Kiến trúc đầu cuối và phần cứng được thay bằng simulator')
para(d,'API xác thực trước khi xử lý; logic chung ra quyết định; SQLite lưu bền vững; web đọc theo quyền, poll 5 giây và nhận SSE thông báo thay đổi. Cloud là phương án tương lai, bản demo chạy cục bộ.')
d.add_page_break()

heading(d,'2 Dữ liệu và giao tiếp giữa các thành phần')
image(d,'erd')
caption(d,'Hình 2  Các thực thể cốt lõi và quan hệ')
para(d,'Users và site_users xác định quyền manager/staff. Devices lưu định danh, trạng thái, thời điểm và bộ đếm. Readings lưu lịch sử; unique (device_id,message_id) chống trùng. Alert_rules có version. Alerts dùng episode_id cho một đợt, tách ack_by/ack_at khỏi resolved_at; audit_logs truy vết thao tác. Incidents/notes hỗ trợ quy trình xử lý.')
table(d,['API cốt lõi','Dữ liệu và quyền'],[
 ('POST /api/telemetry','X-Device-Key; device_id, message_id, measured_at UTC, level_cm, quality'),
 ('POST /api/auth/login; GET /api/auth/me','Email/mật khẩu; phiên và quyền theo cơ sở'),
 ('GET /api/sites/:id/overview','Mực nước, thời điểm, data_status, mức cảnh báo; thành viên'),
 ('GET /api/devices/:id/readings','Lịch sử; kiểm tra cơ sở của thiết bị'),
 ('POST /api/alerts/:id/ack','Thành viên; ghi tiếp nhận, không kết thúc'),
 ('PATCH /api/devices/:id/rule','Manager; ngưỡng, sample_count, expected_version')],[7.3,10.0])
para(d,'Bản tin có device_id, message_id dạng UUID, measured_at UTC, level_cm và quality. Ví dụ mức nước 12 cm với quality="ok"; lỗi cảm biến có null. Timestamp sai, quality lạ hoặc số đo ngoài miền bị từ chối. Mẫu định dạng đầy đủ nằm trong docs/api.md.')
para(d,'Cùng ID/cùng payload trả duplicate, không đếm lại; cùng ID/khác payload trả 409. Cũ quá 15 giây hoặc không mới hơn hiện trạng chỉ lưu history_only, không đổi cảnh báo. Reading và state được ghi trong giao dịch; thông báo sau COMMIT.')
d.add_page_break()

heading(d,'3 Luồng hoạt động và phần thiết kế bổ sung TH4')
image(d,'sequence-telemetry',17.0)
caption(d,'Hình 3  Sequence nhận số đo mới, lưu trữ và cập nhật website')
image(d,'sequence-ack',17.0)
caption(d,'Hình 4  Sequence xác nhận tiếp nhận và kết thúc bởi dữ liệu phục hồi')
para(d,'Bổ sung TH4: logic cảnh báo và vận hành khi dữ liệu gián đoạn. Ba mẫu tươi liên tiếp ≥10 cm mở warning; ≥20 cm mở/nâng critical; <7 cm mới đóng đợt. Không tự hạ critical về warning. Ack chỉ xác nhận con người đã tiếp nhận; không đồng nghĩa nước đã rút.')
para(d,'Mẫu lỗi hoặc gap >30 giây reset đếm; mẫu cũ/trùng không tăng đếm. Không nhận bản tin hợp lệ >30 giây là offline; chỉ nhận dữ liệu cũ là stale. Hai trạng thái này không tự trả normal. Đổi luật có version, reset đếm, ghi audit. Các quyết định và lưu đồ chi tiết nằm trong phiếu TH4 đính kèm.')
d.add_page_break()

heading(d,'4 Bảo mật và minh chứng kiểm thử')
para(d,'Threat Model bảo vệ danh tính, khóa thiết bị, số đo, ngưỡng, cảnh báo và lịch sử thao tác qua ba ranh giới: simulator–API, web–API, API–DB. Cảnh báo hỗ trợ người vận hành; chưa phải hệ thống bảo đảm an toàn thực địa.')
table(d,['Rủi ro','Biện pháp thiết kế'],[
 ('Giả mạo thiết bị, lộ khóa','Khóa riêng, lưu hash; chỉ cấp khóa gốc một lần; API không trả hash'),
 ('Truy cập sai vai trò/cơ sở','Session và kiểm tra role/site phía server cho từng tài nguyên'),
 ('Replay, mẫu cũ, ghi đồng thời','Unique ID, history_only, transaction tuần tự và rollback'),
 ('CSRF, chiếm phiên','Origin-check, JSON, SameSite=Lax, HttpOnly; Secure khi HTTPS'),
 ('Injection/XSS và quá tải','SQL tham số hóa, render text an toàn, limit payload/rate'),
 ('Chối bỏ thay đổi','Audit người, thời điểm, đối tượng và giá trị; không log secret')],[5.0,12.3])
table(d,['Ca bảo mật','Điều kiện và kết quả mong đợi'],[
 ('S01 Phiên','Chưa login/sau logout gọi API bảo vệ → 401'),
 ('S02 Khóa thiết bị','Thiếu/sai khóa → 401; khóa A dùng cho B → 403; không thêm reading'),
 ('S03 Phân quyền','Staff sửa ngưỡng hoặc đọc ID ngoài cơ sở → chặn, DB giữ nguyên'),
 ('S04 Đầu vào','Quality/time/số đo sai → 400, không đổi state'),
 ('S05 Replay','Retry không thêm; khác payload cùng ID → 409; mẫu cũ chỉ lịch sử'),
 ('S06 CSRF','Origin khác gửi thao tác đổi dữ liệu → chặn trước DB'),
 ('S07 Injection/XSS','Không thực thi SQL/script; JSON quá lớn bị từ chối'),
 ('S08 Rò rỉ khóa','Response/log/Git không có khóa gốc hoặc hash')],[4.2,13.1])
summary_path=HERE/'test-summary.txt'
summary=summary_path.read_text(encoding='utf-8').strip() if summary_path.exists() else 'Kết quả chạy kiểm thử đang được đối chiếu với log thực; các dòng trên là tiêu chí, chưa tự coi là PASS. Xem docs/test-results.md để cập nhật trước nộp.'
para(d,summary,True)
para(d,'Các kiểm thử chức năng bao gồm biên 10/20/7 cm và ba mẫu, ack không resolve, không lặp alert, lỗi/gap, stale/offline, đổi luật, restart và concurrency. Chưa có phép thử cảm biến, Wi-Fi, pin hoặc hiện trường.')
d.add_page_break()

heading(d,'5 Thương mại hóa và tổ chức thực hiện')
para(d,'Khách hàng dự kiến là chủ bãi xe nhỏ/đơn vị vận hành khu vực trũng; người dùng là quản lý và nhân viên trực. Giá trị là theo dõi tập trung, phát hiện vượt ngưỡng và phối hợp tiếp nhận. Đề xuất bán gói ba điểm đo kèm lắp đặt, thu phí web theo cơ sở/tháng. Chưa có khảo sát khách hàng hay xác nhận sẵn sàng trả tiền.')
table(d,['Khoản mục mỗi điểm đo','Chi phí giả định nghìn đồng'],[
 ('Cảm biến siêu âm','350'),('ESP32 và Wi-Fi','150'),('Nguồn và dự phòng','200'),
 ('Hộp, giá đỡ, phụ kiện','150'),('Lắp đặt và hiệu chuẩn','250'),('Dự phòng bảo hành','100'),
 ('Tổng chi phí trực tiếp','1.200'),('Giá bán đề xuất / phí web cơ sở mỗi tháng','1.800 / 99')],[10.5,6.8])
para(d,'Tất cả giá là giả định, không phải báo giá. Gói ba điểm: chi phí trực tiếp 3,6 triệu đồng, giá đề xuất 5,4 triệu đồng; chênh lệch 1,8 triệu chưa trừ phần mềm, hỗ trợ, marketing, thuế và vận hành, không phải lợi nhuận ròng.')
para(d,'Bước tiếp theo: phỏng vấn 5–10 cơ sở, xin báo giá, thử một điểm đối chiếu thước đo và Wi-Fi; hiệu chỉnh vị trí/ngưỡng/giá trước triển khai. Không có bơm tự động, SMS/email, dự báo hoặc thiết bị thật trong phạm vi demo.')
para(d,'Phân công đề xuất để nhóm xác nhận',True)
table(d,['Thành viên và MSSV','Phần việc đề xuất'],[
 ('Hà Minh Tiến 1250080194','Backend, DB/API, logic và simulator'),
 ('Trương Quốc Toàn 1250080203','Frontend, dashboard, lịch sử và trải nghiệm'),
 ('Huỳnh Nguyễn Duy Tiến 1250080195','Kiểm thử, Threat Model, tài liệu/slide/demo')],[8.1,9.2])
para(d,'Repository: https://github.com/tienpp1/TMDT. Source, README, sơ đồ, API, Threat Model, kịch bản demo và phiếu TH4 phục vụ đối chiếu/bảo vệ. Báo cáo là tài liệu nộp theo hướng dẫn học phần; lớp chưa được cung cấp.')
para(d,'Nguồn: Do_an_cuoi_ky_TMDT.docx (Hướng 2, nội dung tối thiểu, demo); TH4 Bổ sung thiết kế cho đồ án cuối kỳ.docx (thiết kế và phiếu). Nhóm cần hiểu, kiểm tra và tự chỉnh bài theo quy định AI của học phần.')
d.save(HERE/'Bao_cao_IoT_Nhom10.docx')

# Patch only document.xml of the retained TH4; all other ZIP parts byte-preserved.
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W='{'+NS['w']+'}'
ET.register_namespace('w',NS['w'])
def alltext(el): return ''.join(n.text or '' for n in el.iter(W+'t'))
with ZipFile(SOURCE) as zin:
    raw=zin.read('word/document.xml'); root=ET.fromstring(raw)
    for p in root.findall('.//w:body/w:p',NS):
        if alltext(p).strip().startswith('Nhóm:'):
            ts=p.findall('.//w:t',NS)
            if ts:
                ts[0].text='Nhóm: 10'
                for t in ts[1:]: t.text=''
    tb=root.findall('.//w:body/w:tbl',NS)[-1]
    values=[
      ['Thiết kế và mô phỏng hệ thống IoT giám sát mực nước và cảnh báo ngập cho bãi giữ xe qua website.'],
      ['☐ 1 — Hệ thống thương mại điện tử     ☑ 2 — Sản phẩm IoT'],
      ['Logic cảnh báo và vận hành khi dữ liệu gián đoạn.'],
      ['Quản lý và nhân viên trực của một bãi xe; ba điểm đo giả lập. Theo dõi nước dâng, tiếp nhận cảnh báo khi Wi-Fi/số đo có thể gián đoạn.'],
      ['Tránh cảnh báo lặp do dao động, mẫu cũ ghi đè hiện trạng và hiểu nhầm ack/mất mạng là đã an toàn. Bổ sung chuỗi ba mẫu, độ mới và xác nhận độc lập để xử lý nhất quán.'],
      ['☑ Mô hình dữ liệu  ☑ Sơ đồ kiến trúc  ☑ Lưu đồ logic  ☑ Khác: hai Sequence Diagram.','Đính kèm figures/architecture, erd, state-alerts, sequence-telemetry, sequence-ack (SVG/PNG) và báo cáo Nhóm 10.'],
      ['1. Ba mẫu tươi liên tiếp: ≥10 cm warning; ≥20 cm critical; <7 cm kết thúc. Ưu tiên critical; ngưỡng demo cần hiệu chỉnh khi thử thật.',
       '2. Ack ghi người/thời điểm tiếp nhận, không resolve. Offline/stale/sensor_error không tự đổi cảnh báo về normal.',
       '3. Message_id unique chống retry; cùng ID khác payload trả 409. Cũ >15 giây hoặc không mới hơn hiện trạng chỉ lưu history_only.',
       '4. Lỗi cảm biến có level_cm=null; gap >30 giây reset đếm. Không nhận bản tin hợp lệ >30 giây là offline; chỉ nhận dữ liệu cũ là stale.',
       '5. Simulator thay cảm biến/ESP32, gửi HTTP mỗi 5 giây; backend quyết định chung và lưu SQLite; web poll 5 giây/SSE. Manager đổi luật có version, reset đếm và audit.'],
      ['Thành phần thêm/sửa: logic telemetry dùng chung, version luật, độ mới dashboard, kịch bản simulator và test.',
       'Dữ liệu liên quan: devices, readings, alert_rules, alerts, audit_logs; API telemetry/tổng quan/lịch sử. Đầu đo, lỗi và ngắt mạng được mô phỏng, không bắt buộc chế tạo phần cứng.'],
      ['Đối chiếu log test/demo; nhóm xác nhận phân công; hiệu chỉnh vị trí/ngưỡng với người trực. Sau đó khảo sát khách hàng, xin báo giá và thử cảm biến/nguồn/Wi-Fi khi có phần cứng. Chưa có kết quả thực địa.']
    ]
    for row,lines in zip(tb.findall('w:tr',NS),values):
        cell=row.findall('w:tc',NS)[1]
        sample=cell.find('w:p',NS)
        sample_ppr=deepcopy(sample.find('w:pPr',NS)) if sample is not None and sample.find('w:pPr',NS) is not None else None
        sample_rpr=deepcopy(sample.find('.//w:rPr',NS)) if sample is not None and sample.find('.//w:rPr',NS) is not None else None
        for el in list(cell):
            if el.tag!=W+'tcPr': cell.remove(el)
        for line in lines:
            p=ET.SubElement(cell,W+'p')
            if sample_ppr is not None: p.append(deepcopy(sample_ppr))
            r=ET.SubElement(p,W+'r')
            if sample_rpr is not None: r.append(deepcopy(sample_rpr))
            ET.SubElement(r,W+'t').text=line
    target=HERE/'TH4_hoan_thanh.docx'
    with ZipFile(target,'w',ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(deepcopy(item),ET.tostring(root,encoding='utf-8',xml_declaration=True) if item.filename=='word/document.xml' else zin.read(item.filename))
    inventory={i.filename:sha256(zin.read(i.filename)).hexdigest() for i in zin.infolist()}
    with ZipFile(target) as zout:
        changed=[n for n,h in inventory.items() if sha256(zout.read(n)).hexdigest()!=h]
    contract={
      'source':str(SOURCE),'source_sha256':sha256(SOURCE.read_bytes()).hexdigest(),
      'changed_parts':changed,'preserve_only_parts_unchanged':changed==['word/document.xml'],
      'source_tables':len(ET.fromstring(raw).findall('.//w:tbl',NS)),
      'final_tables':len(root.findall('.//w:tbl',NS)),
      'edit_scope':'Nhóm paragraph and right-hand cells of final completion form only',
      'package_inventory':inventory,
    }
    (HERE/'template-preservation.json').write_text(json.dumps(contract,ensure_ascii=False,indent=2),encoding='utf-8')
    sect=root.find('.//w:sectPr',NS)
    (HERE/'artifact.md').write_text('# Hợp đồng điền phiếu TH4\n\n'
      f'Nguồn giữ nguyên: {SOURCE}\n\nSHA256: {contract["source_sha256"]}\n\n'
      'Chỉ sửa word/document.xml: đoạn Nhóm và ô phải bảng cuối. Giữ nguyên các phần yêu cầu, nhãn hàng, phần XML khác, styles, numbering, header/footer và relationships. '
      'Tất cả parts còn lại so byte SHA256 trong template-preservation.json. Không dựng lại template từ tài liệu trắng. '
      'Giữ page geometry, style và table/cell properties từ nguồn; paragraph/run properties được sao chép từ chính ô đang điền. '
      'Nội dung dài được thêm paragraph theo mẫu của ô; không giảm cỡ chữ. Số trang và mọi trang cần render QA.\n\n'
      'Section XML: '+ET.tostring(sect,encoding='unicode')+'\n',encoding='utf-8')
print('Created report, filled TH4, five SVG/PNG figures, preservation evidence.')
